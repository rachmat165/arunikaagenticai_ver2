"""
Membaca dan mengekstrak teks dari berbagai format file attachment Telegram.
Mendukung: PDF, DOCX, XLSX, PPTX, TXT, MD, CSV, JSON, XML, HTML, PY,
           dan gambar/screenshot (JPG, PNG, WebP, GIF, BMP, TIFF) via AI Vision.
"""

import asyncio
from pathlib import Path


SUPPORTED_TEXT = {
    ".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".java", ".cpp", ".c", ".h",
    ".yaml", ".yml", ".toml", ".ini", ".cfg", ".env",
    ".log", ".sql", ".sh", ".bat", ".ps1",
}
SUPPORTED_PDF  = {".pdf"}
SUPPORTED_DOCX = {".docx", ".doc"}
SUPPORTED_XLSX = {".xlsx", ".xls"}
SUPPORTED_PPTX = {".pptx", ".ppt"}
SUPPORTED_IMG  = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp", ".tiff", ".tif", ".heic"}

MAX_CHARS = 12_000

ALL_SUPPORTED = (
    SUPPORTED_TEXT | SUPPORTED_PDF | SUPPORTED_DOCX |
    SUPPORTED_XLSX | SUPPORTED_PPTX
)

FORMAT_LABEL = (
    "PDF, DOCX, XLSX, PPTX, TXT, MD, CSV, JSON, XML, HTML, PY, "
    "JS, TS, YAML, LOG, SQL, SH, BAT\n"
    "📸 Gambar: JPG, PNG, WebP, GIF, BMP (kirim sebagai foto atau file)"
)


async def extract_text(file_path: str, file_name: str) -> dict:
    """
    Ekstrak teks dari file.
    Returns {'text': str, 'pages': int, 'file_name': str, 'file_type': str}
    atau {'error': str}
    """
    ext = Path(file_name).suffix.lower()

    def _read():
        if ext in SUPPORTED_TEXT:
            return _read_text(file_path)
        elif ext in SUPPORTED_PDF:
            return _read_pdf(file_path)
        elif ext in SUPPORTED_DOCX:
            return _read_docx(file_path)
        elif ext in SUPPORTED_XLSX:
            return _read_xlsx(file_path)
        elif ext in SUPPORTED_PPTX:
            return _read_pptx(file_path)
        else:
            return {
                "error": (
                    f"Format *{ext.upper()}* belum didukung untuk ekstraksi teks.\n\n"
                    f"Format yang bisa dibaca:\n{FORMAT_LABEL}"
                )
            }

    try:
        result = await asyncio.to_thread(_read)
        if "error" not in result:
            result["file_name"] = file_name
            result["file_type"] = ext.lstrip(".")
        return result
    except Exception as e:
        return {"error": f"Gagal membaca file: {e}"}


def _read_text(path: str) -> dict:
    for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = Path(path).read_text(encoding=enc)
            return {"text": text[:MAX_CHARS], "pages": 1}
        except UnicodeDecodeError:
            continue
    return {"error": "Tidak bisa decode file teks (encoding tidak dikenal)"}


def _read_pdf(path: str) -> dict:
    try:
        import pdfplumber
        parts = []
        pages = 0
        with pdfplumber.open(path) as pdf:
            pages = len(pdf.pages)
            for page in pdf.pages[:40]:
                t = page.extract_text()
                if t:
                    parts.append(t.strip())
        text = "\n\n".join(parts)
        if not text.strip():
            return {
                "error": (
                    "PDF tidak mengandung teks yang bisa diekstrak.\n"
                    "Kemungkinan PDF ini berupa scan/gambar.\n"
                    "💡 Kirim sebagai *foto* agar dianalisis dengan Vision AI."
                )
            }
        return {"text": text[:MAX_CHARS], "pages": pages}
    except ImportError:
        return {"error": "Library pdfplumber tidak tersedia. Jalankan: pip install pdfplumber"}
    except Exception as e:
        return {"error": f"Gagal membaca PDF: {e}"}


def _read_docx(path: str) -> dict:
    try:
        from docx import Document
        doc = Document(path)
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())
        # Baca juga tabel
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    parts.append(row_text)
        text = "\n".join(parts)
        if not text.strip():
            return {"error": "Dokumen DOCX kosong atau tidak mengandung teks"}
        return {"text": text[:MAX_CHARS], "pages": 1}
    except ImportError:
        return {"error": "Library python-docx tidak tersedia. Jalankan: pip install python-docx"}
    except Exception as e:
        return {"error": f"Gagal membaca DOCX: {e}"}


def _read_xlsx(path: str) -> dict:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            parts.append(f"=== Sheet: {sheet.title} ===")
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(c) if c is not None else "" for c in row]
                line = " | ".join(row_vals).strip(" |")
                if line:
                    parts.append(line)
        wb.close()
        text = "\n".join(parts)
        if not text.strip():
            return {"error": "File Excel kosong"}
        return {"text": text[:MAX_CHARS], "pages": len(wb.worksheets)}
    except ImportError:
        return {"error": "Library openpyxl tidak tersedia. Jalankan: pip install openpyxl"}
    except Exception as e:
        return {"error": f"Gagal membaca XLSX: {e}"}


def _read_pptx(path: str) -> dict:
    try:
        from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                parts.append(f"=== Slide {i} ===\n" + "\n".join(slide_texts))
        text = "\n\n".join(parts)
        if not text.strip():
            return {"error": "Presentasi PPTX tidak mengandung teks"}
        return {"text": text[:MAX_CHARS], "pages": len(prs.slides)}
    except ImportError:
        return {"error": "Library python-pptx tidak tersedia. Jalankan: pip install python-pptx"}
    except Exception as e:
        return {"error": f"Gagal membaca PPTX: {e}"}


def is_image(file_name: str) -> bool:
    """Cek apakah file adalah gambar/screenshot yang perlu diproses via Vision AI."""
    return Path(file_name).suffix.lower() in SUPPORTED_IMG


def is_supported(file_name: str) -> bool:
    """Cek apakah file bisa diekstrak teksnya."""
    return Path(file_name).suffix.lower() in ALL_SUPPORTED
